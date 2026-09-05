# PowerPoint text extraction.

from pptx import Presentation
from .base_extractor import create_document_template
import os
from PIL import Image
from io import BytesIO
from pptx.enum.shapes import MSO_SHAPE_TYPE
from zipfile import ZipFile
import shutil
from lxml import etree
from .output_paths import extracted_dir


def extract_metadata(presentation):

    props = presentation.core_properties

    metadata = {

        "author": props.author,

        "category": props.category,

        "comments": props.comments,

        "content_status": props.content_status,

        "created": str(props.created) if props.created else "",

        "identifier": props.identifier,

        "keywords": props.keywords,

        "language": props.language,

        "last_modified_by": props.last_modified_by,

        "last_printed": str(props.last_printed) if props.last_printed else "",

        "modified": str(props.modified) if props.modified else "",

        "revision": props.revision,

        "subject": props.subject,

        "title": props.title,

        "version": props.version

    }

    return metadata


def extract_slides(presentation):

    slides = []

    for slide_number, slide in enumerate(presentation.slides, start=1):

        slide_info = {

            "slide_number": slide_number,

            "shape_count": len(slide.shapes),

            "layout": slide.slide_layout.name

        }

        slides.append(slide_info)

    return slides

def extract_titles(presentation):

    titles = []

    for slide_number, slide in enumerate(presentation.slides, start=1):

        title_text = ""

        if slide.shapes.title:
            title_text = slide.shapes.title.text.strip()

        titles.append({

            "slide_number": slide_number,

            "title": title_text

        })

    return titles


def extract_text(presentation):

    text_data = []

    for slide_number, slide in enumerate(presentation.slides, start=1):

        for shape_number, shape in enumerate(slide.shapes, start=1):

            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            if text:

                text_data.append({

                    "slide_number": slide_number,

                    "shape_number": shape_number,

                    "text": text

                })

    return text_data

def extract_tables(presentation):

    tables = []

    table_number = 1

    for slide_number, slide in enumerate(presentation.slides, start=1):

        for shape in slide.shapes:

            if not shape.has_table:
                continue

            table = shape.table

            table_data = []

            for row in table.rows:

                row_data = []

                for cell in row.cells:

                    row_data.append(cell.text.strip())

                table_data.append(row_data)

            tables.append({

                "table_number": table_number,

                "slide_number": slide_number,

                "rows": len(table.rows),

                "columns": len(table.columns),

                "data": table_data

            })

            table_number += 1

    return tables

def extract_images(presentation, file_path):

    images = []

    image_output_dir = str(extracted_dir("images", file_path))

    os.makedirs(image_output_dir, exist_ok=True)

    image_number = 1

    for slide_number, slide in enumerate(presentation.slides, start=1):

        for shape in slide.shapes:

            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue

            image = shape.image

            extension = image.ext

            image_name = f"image_{image_number}.{extension}"

            image_path = os.path.join(image_output_dir, image_name)

            with open(image_path, "wb") as f:
                f.write(image.blob)

            try:

                img = Image.open(BytesIO(image.blob))

                width, height = img.size

            except Exception:

                width = None
                height = None

            images.append({

                "image_number": image_number,

                "slide_number": slide_number,

                "image_name": image_name,

                "image_path": image_path,

                "width": width,

                "height": height

            })

            image_number += 1

    return images

def extract_charts(presentation):

    charts = []

    chart_number = 1

    for slide_number, slide in enumerate(presentation.slides, start=1):

        for shape in slide.shapes:

            if not shape.has_chart:
                continue

            chart = shape.chart

            chart_info = {

                "chart_number": chart_number,

                "slide_number": slide_number,

                "chart_type": str(chart.chart_type),

                "categories": [],

                "series": []

            }

            try:

                categories = chart.plots[0].categories

                chart_info["categories"] = [
                    str(category.label)
                    for category in categories
                ]

            except Exception:
                pass

            try:

                for series in chart.series:

                    values = []

                    try:
                        values = list(series.values)
                    except Exception:
                        pass

                    chart_info["series"].append({

                        "name": series.name,

                        "values": values

                    })

            except Exception:
                pass

            charts.append(chart_info)

            chart_number += 1

    return charts

def extract_hyperlinks(presentation):

    hyperlinks = []

    hyperlink_number = 1

    for slide_number, slide in enumerate(presentation.slides, start=1):

        for shape_number, shape in enumerate(slide.shapes, start=1):

            try:

                if shape.click_action.hyperlink.address:

                    hyperlinks.append({

                        "hyperlink_number": hyperlink_number,

                        "slide_number": slide_number,

                        "shape_number": shape_number,

                        "text": shape.text if hasattr(shape, "text") else "",

                        "url": shape.click_action.hyperlink.address

                    })

                    hyperlink_number += 1

            except Exception:
                continue

    return hyperlinks

def extract_notes(presentation):

    notes = []

    for slide_number, slide in enumerate(presentation.slides, start=1):

        note_text = ""

        try:

            notes_slide = slide.notes_slide

            if notes_slide:

                for shape in notes_slide.shapes:

                    if hasattr(shape, "text"):

                        text = shape.text.strip()

                        if text:

                            note_text += text + "\n"

        except Exception:
            pass

        notes.append({

            "slide_number": slide_number,

            "notes": note_text.strip()

        })

    return notes
def extract_shapes(presentation):

    shapes = []

    shape_number = 1

    for slide_number, slide in enumerate(presentation.slides, start=1):

        for shape in slide.shapes:

            try:
                shape_text = shape.text.strip() if hasattr(shape, "text") else ""
            except Exception:
                shape_text = ""

            shapes.append({

                "shape_number": shape_number,

                "slide_number": slide_number,

                "name": shape.name,

                "shape_type": str(shape.shape_type),

                "left": shape.left,

                "top": shape.top,

                "width": shape.width,

                "height": shape.height,

                "has_text": hasattr(shape, "text"),

                "text": shape_text

            })

            shape_number += 1

    return shapes

def extract_smartart(file_path):
    """
    Extract text from SmartArt diagrams in a PowerPoint presentation.
    """

    smartart = []

    try:

        with ZipFile(file_path) as ppt:

            smartart_number = 1

            ns = {
                "a": "http://schemas.openxmlformats.org/drawingml/2006/main"
            }

            # SmartArt data is usually stored in ppt/diagrams/
            for file in ppt.namelist():

                if not file.startswith("ppt/diagrams/"):
                    continue

                if not file.endswith(".xml"):
                    continue

                try:

                    xml = ppt.read(file)

                    root = etree.fromstring(xml)

                    texts = []

                    for node in root.findall(".//a:t", ns):

                        if node.text:
                            texts.append(node.text)

                    smartart.append({

                        "diagram_number": smartart_number,

                        "diagram_file": os.path.basename(file),

                        "text": texts

                    })

                    smartart_number += 1

                except Exception:
                    continue

    except Exception as e:

        smartart.append({

            "error": str(e)

        })

    return smartart

def extract_embedded_objects(file_path):

    embedded_objects = []

    output_folder = str(extracted_dir("embedded", file_path))

    os.makedirs(output_folder, exist_ok=True)

    try:

        with ZipFile(file_path) as ppt:

            files = ppt.namelist()

            object_number = 1

            for file in files:

                if not file.startswith("ppt/embeddings/"):
                    continue

                filename = os.path.basename(file)

                destination = os.path.join(output_folder, filename)

                with ppt.open(file) as source, open(destination, "wb") as target:
                    shutil.copyfileobj(source, target)

                embedded_objects.append({

                    "object_number": object_number,

                    "file_name": filename,

                    "file_path": destination

                })

                object_number += 1

    except Exception as e:

        embedded_objects.append({

            "error": str(e)

        })

    return embedded_objects

def extract_videos(file_path):

    videos = []

    output_folder = str(extracted_dir("videos", file_path))

    os.makedirs(output_folder, exist_ok=True)

    video_extensions = (
        ".mp4",
        ".avi",
        ".wmv",
        ".mov",
        ".mkv",
        ".mpeg",
        ".mpg"
    )

    try:

        with ZipFile(file_path) as ppt:

            video_number = 1

            for file in ppt.namelist():

                if not file.startswith("ppt/media/"):
                    continue

                filename = os.path.basename(file)

                if not filename.lower().endswith(video_extensions):
                    continue

                destination = os.path.join(output_folder, filename)

                with ppt.open(file) as source, open(destination, "wb") as target:
                    shutil.copyfileobj(source, target)

                videos.append({

                    "video_number": video_number,

                    "video_name": filename,

                    "video_path": destination

                })

                video_number += 1

    except Exception as e:

        videos.append({

            "error": str(e)

        })

    return videos

def extract_audio(file_path):

    audio = []

    output_folder = str(extracted_dir("audio", file_path))

    os.makedirs(output_folder, exist_ok=True)

    audio_extensions = (
        ".mp3",
        ".wav",
        ".aac",
        ".m4a",
        ".wma",
        ".flac",
        ".ogg"
    )

    try:

        with ZipFile(file_path) as ppt:

            audio_number = 1

            for file in ppt.namelist():

                if not file.startswith("ppt/media/"):
                    continue

                filename = os.path.basename(file)

                if not filename.lower().endswith(audio_extensions):
                    continue

                destination = os.path.join(output_folder, filename)

                with ppt.open(file) as source, open(destination, "wb") as target:
                    shutil.copyfileobj(source, target)

                audio.append({

                    "audio_number": audio_number,

                    "audio_name": filename,

                    "audio_path": destination

                })

                audio_number += 1

    except Exception as e:

        audio.append({

            "error": str(e)

        })

    return audio

def extract(file_path):

    presentation = Presentation(file_path)

    result = create_document_template(file_path)

    result["document"]["file_type"] = "powerpoint"

    result["metadata"] = extract_metadata(presentation)

    result["slides"] = extract_slides(presentation)

    result["titles"] = extract_titles(presentation)

    result["text"] = extract_text(presentation)

    result["tables"] = extract_tables(presentation)

    result["images"] = extract_images(presentation, file_path)

    result["charts"] = extract_charts(presentation)

    result["hyperlinks"] = extract_hyperlinks(presentation)

    result["notes"] = extract_notes(presentation)

    result["shapes"] = extract_shapes(presentation)

    result["smartart"] = extract_smartart(file_path)

    result["embedded_objects"] = extract_embedded_objects(file_path)

    result["videos"] = extract_videos(file_path)

    result["audio"] = extract_audio(file_path)

    return result
