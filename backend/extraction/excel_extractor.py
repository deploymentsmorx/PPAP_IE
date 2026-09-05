# Excel workbook text extraction.

from openpyxl import load_workbook
import os
from .base_extractor import create_document_template
from .output_paths import extracted_dir
import win32com.client


# -------------------------------------------------
# Metadata
# -------------------------------------------------

def extract_metadata(workbook):
    """
    Extract workbook metadata.
    """

    props = workbook.properties

    return {

        "creator": props.creator,

        "last_modified_by": props.lastModifiedBy,

        "created": str(props.created) if props.created else "",

        "modified": str(props.modified) if props.modified else "",

        "title": props.title,

        "subject": props.subject,

        "description": props.description,

        "keywords": props.keywords,

        "category": props.category,

        "company": getattr(props, "company", "")

    }


# -------------------------------------------------
# Worksheets
# -------------------------------------------------

def extract_worksheets(workbook):
    """
    Extract worksheet details.
    """

    worksheets = []

    for sheet in workbook.worksheets:

        worksheets.append({

            "sheet_name": sheet.title,

            "max_rows": sheet.max_row,

            "max_columns": sheet.max_column,

            "sheet_state": sheet.sheet_state

        })

    return worksheets

def extract_tables(workbook):
    """
    Extract worksheet data as structured tables.
    """

    tables = []

    table_number = 1

    for sheet in workbook.worksheets:

        rows = list(sheet.iter_rows(values_only=True))

        if not rows:
            continue

        cleaned_rows = []

        for row in rows:

            # Skip completely empty rows
            if row is None:
                continue

            cleaned = []

            is_empty = True

            for cell in row:

                if cell is None:
                    cleaned.append("")
                else:
                    cleaned.append(str(cell))
                    is_empty = False

            if not is_empty:
                cleaned_rows.append(cleaned)

        if not cleaned_rows:
            continue

        header = cleaned_rows[0]

        data_rows = cleaned_rows[1:]

        tables.append({

            "sheet": sheet.title,

            "table_number": table_number,

            "header": header,

            "row_count": len(data_rows),

            "column_count": len(header),

            "rows": data_rows

        })

        table_number += 1

    return tables

def extract_images(workbook, file_path):
    """
    Extract embedded images from Excel worksheets.
    """

    images = []

    image_output_dir = str(extracted_dir("images", file_path))

    os.makedirs(image_output_dir, exist_ok=True)

    image_number = 1

    for sheet in workbook.worksheets:

        if not hasattr(sheet, "_images"):
            continue

        for image in sheet._images:

            image_name = f"image_{image_number}.png"

            image_path = os.path.join(image_output_dir, image_name)

            try:
                # Save embedded image
                with open(image_path, "wb") as img_file:
                    img_file.write(image._data())

                images.append({

                    "sheet": sheet.title,

                    "image_number": image_number,

                    "image_name": image_name,

                    "image_path": image_path

                })

                image_number += 1

            except Exception as e:

                images.append({

                    "sheet": sheet.title,

                    "error": str(e)

                })

    return images

def extract_charts(workbook):
    """
    Extract chart information from all worksheets.
    """

    charts = []

    chart_number = 1

    for sheet in workbook.worksheets:

        if not hasattr(sheet, "_charts"):
            continue

        for chart in sheet._charts:

            charts.append({

                "sheet": sheet.title,

                "chart_number": chart_number,

                "chart_type": type(chart).__name__,

                "title": chart.title.tx.rich.p[0].r[0].t
                if chart.title and hasattr(chart.title, "tx")
                else ""

            })

            chart_number += 1

    return charts

def extract_merged_cells(workbook):
    """
    Extract all merged cell ranges from every worksheet.
    """

    merged_cells = []

    for sheet in workbook.worksheets:

        for merged_range in sheet.merged_cells.ranges:

            merged_cells.append({

                "sheet": sheet.title,

                "range": str(merged_range)

            })

    return merged_cells

def extract_formulas(workbook):
    """
    Extract all formulas from every worksheet.
    """

    formulas = []

    for sheet in workbook.worksheets:

        for row in sheet.iter_rows():

            for cell in row:

                if cell.data_type == "f":

                    formulas.append({

                        "sheet": sheet.title,

                        "cell": cell.coordinate,

                        "formula": cell.value

                    })

    return formulas

def extract_hyperlinks(workbook):
    """
    Extract hyperlinks from all worksheets.
    """

    hyperlinks = []

    for sheet in workbook.worksheets:

        for row in sheet.iter_rows():

            for cell in row:

                if cell.hyperlink:

                    hyperlinks.append({

                        "sheet": sheet.title,

                        "cell": cell.coordinate,

                        "text": cell.value,

                        "url": cell.hyperlink.target

                    })

    return hyperlinks

def extract_comments(workbook):
    """
    Extract comments from all worksheets.
    """

    comments = []

    for sheet in workbook.worksheets:

        for row in sheet.iter_rows():

            for cell in row:

                if cell.comment:

                    comments.append({

                        "sheet": sheet.title,

                        "cell": cell.coordinate,

                        "author": cell.comment.author,

                        "text": cell.comment.text

                    })

    return comments

def extract_named_ranges(workbook):
    """
    Extract all named ranges from the workbook.
    """

    named_ranges = []

    try:

        for defined_name in workbook.defined_names.values():

            named_ranges.append({

                "name": defined_name.name,

                "value": defined_name.attr_text

            })

    except Exception as e:

        named_ranges.append({

            "error": str(e)

        })

    return named_ranges

def extract_pivot_tables(file_path):
    """
    Extract Pivot Tables using Microsoft Excel COM.
    """

    pivot_tables = []

    excel = None
    workbook = None

    try:

        excel = win32com.client.Dispatch("Excel.Application")

        excel.Visible = False

        workbook = excel.Workbooks.Open(file_path)

        pivot_number = 1

        for worksheet in workbook.Worksheets:

            for pivot in worksheet.PivotTables():

                pivot_info = {

                    "sheet": worksheet.Name,

                    "pivot_number": pivot_number,

                    "pivot_name": pivot.Name,

                    "source_data": pivot.SourceData,

                    "row_fields": [],

                    "column_fields": [],

                    "data_fields": [],

                    "page_fields": []

                }

                # Row Fields
                for field in pivot.RowFields():

                    pivot_info["row_fields"].append(field.Name)

                # Column Fields
                for field in pivot.ColumnFields():

                    pivot_info["column_fields"].append(field.Name)

                # Data Fields
                for field in pivot.DataFields():

                    pivot_info["data_fields"].append(field.Name)

                # Filter Fields
                for field in pivot.PageFields():

                    pivot_info["page_fields"].append(field.Name)

                pivot_tables.append(pivot_info)

                pivot_number += 1

    except Exception as e:

        pivot_tables.append({

            "error": str(e)

        })

    finally:

        if workbook:
            workbook.Close(False)

        if excel:
            excel.Quit()

    return pivot_tables


# -------------------------------------------------
# Main Extractor
# -------------------------------------------------

def extract(file_path):

    workbook = load_workbook(
        file_path,
        data_only=False
    )

    result = create_document_template(file_path)
    result["document"]["file_type"] = "excel"
    result["metadata"] = extract_metadata(workbook)
    result["worksheets"] = extract_worksheets(workbook)
    result["tables"] = extract_tables(workbook)
    result["images"] = extract_images(workbook, file_path)
    result["charts"] = extract_charts(workbook)
    result["merged_cells"] = extract_merged_cells(workbook)
    result["formulas"] = extract_formulas(workbook)
    result["hyperlinks"] = extract_hyperlinks(workbook)
    result["comments"] = extract_comments(workbook)
    result["named_ranges"] = extract_named_ranges(workbook)
    result["pivot_tables"] = extract_pivot_tables(file_path)
    workbook.close()

    return result
