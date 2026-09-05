# Word document text extraction.

import os
from docx import Document
from PIL import Image
from zipfile import ZipFile
from lxml import etree

from .base_extractor import create_document_template
from .output_paths import extracted_dir


# --------------------------------------------------
# Metadata
# --------------------------------------------------

def extract_metadata(document):
    """
    Extract Word document metadata.
    """

    props = document.core_properties

    return {

        "author": props.author,

        "category": props.category,

        "comments": props.comments,

        "content_status": props.content_status,

        "created": str(props.created) if props.created else "",

        "identifier": props.identifier,

        "keywords": props.keywords,

        "language": props.language,

        "last_modified_by": props.last_modified_by,

        "last_printed": str(props.last_printed) if props.last_printed else "",

        "modified": str(props.modified) if props.modified else "",

        "revision": props.revision,

        "subject": props.subject,

        "title": props.title,

        "version": props.version

    }


# --------------------------------------------------
# Paragraphs
# --------------------------------------------------

def extract_paragraphs(document):
    """
    Extract all paragraphs.
    """

    paragraphs = []

    for index, para in enumerate(document.paragraphs, start=1):

        paragraphs.append({

            "paragraph_number": index,

            "style": para.style.name,

            "text": para.text.strip()

        })

    return paragraphs


# --------------------------------------------------
# Headings
# --------------------------------------------------

def extract_headings(document):
    """
    Extract headings.
    """

    headings = []

    for para in document.paragraphs:

        if para.style.name.startswith("Heading"):

            headings.append({

                "style": para.style.name,

                "text": para.text.strip()

            })

    return headings

# --------------------------------------------------
# Tables
# --------------------------------------------------

def extract_tables(document):
    """
    Extract all tables from the Word document.
    """

    tables = []

    for table_number, table in enumerate(document.tables, start=1):

        table_data = []

        for row in table.rows:

            row_data = []

            for cell in row.cells:

                row_data.append(cell.text.strip())

            table_data.append(row_data)

        if table_data:

            tables.append({

                "table_number": table_number,

                "rows": len(table_data),

                "columns": len(table_data[0]) if table_data else 0,

                "data": table_data

            })

    return tables

# --------------------------------------------------
# Images
# --------------------------------------------------

def extract_images(document, file_path):
    """
    Extract embedded images from the Word document.
    """

    images = []

    image_output_dir = str(extracted_dir("images", file_path))

    os.makedirs(image_output_dir, exist_ok=True)

    image_number = 1

    for rel in document.part.rels.values():

        if "image" not in rel.target_ref:
            continue

        image = rel.target_part

        image_name = f"image_{image_number}.{image.content_type.split('/')[-1]}"

        image_path = os.path.join(image_output_dir, image_name)

        with open(image_path, "wb") as f:
            f.write(image.blob)

        try:

            img = Image.open(image_path)

            width, height = img.size

        except Exception:

            width = None
            height = None

        images.append({

            "image_number": image_number,

            "image_name": image_name,

            "image_path": image_path,

            "width": width,

            "height": height

        })

        image_number += 1

    return images

# --------------------------------------------------
# Hyperlinks
# --------------------------------------------------

from docx.oxml.ns import qn


def extract_hyperlinks(document):
    """
    Extract all hyperlinks from the Word document.
    """

    hyperlinks = []

    hyperlink_number = 1

    for paragraph in document.paragraphs:

        for hyperlink in paragraph._p.findall(".//w:hyperlink", paragraph._p.nsmap):

            r_id = hyperlink.get(qn("r:id"))

            if not r_id:
                continue

            try:

                relationship = document.part.rels[r_id]

                url = relationship.target_ref

                text = "".join(
                    node.text
                    for node in hyperlink.findall(".//w:t", paragraph._p.nsmap)
                    if node.text
                )

                hyperlinks.append({

                    "hyperlink_number": hyperlink_number,

                    "text": text,

                    "url": url

                })

                hyperlink_number += 1

            except Exception:
                continue

    return hyperlinks

# --------------------------------------------------
# Headers
# --------------------------------------------------

def extract_headers(document):
    """
    Extract headers from all sections.
    """

    headers = []

    for section_number, section in enumerate(document.sections, start=1):

        for paragraph in section.header.paragraphs:

            text = paragraph.text.strip()

            if text:

                headers.append({

                    "section": section_number,

                    "text": text

                })

        # Header tables
        for table in section.header.tables:

            table_data = []

            for row in table.rows:

                table_data.append(
                    [cell.text.strip() for cell in row.cells]
                )

            headers.append({

                "section": section_number,

                "table": table_data

            })

    return headers

# --------------------------------------------------
# Footers
# --------------------------------------------------

def extract_footers(document):
    """
    Extract footers from all sections.
    """

    footers = []

    for section_number, section in enumerate(document.sections, start=1):

        # Footer paragraphs
        for paragraph in section.footer.paragraphs:

            text = paragraph.text.strip()

            if text:

                footers.append({

                    "section": section_number,

                    "text": text

                })

        # Footer tables
        for table in section.footer.tables:

            table_data = []

            for row in table.rows:

                table_data.append(
                    [cell.text.strip() for cell in row.cells]
                )

            footers.append({

                "section": section_number,

                "table": table_data

            })

    return footers

# --------------------------------------------------
# Sections
# --------------------------------------------------

def extract_sections(document):
    """
    Extract information about document sections.
    """

    sections = []

    for section_number, section in enumerate(document.sections, start=1):

        sections.append({

            "section_number": section_number,

            "orientation": str(section.orientation),

            "page_width": section.page_width,

            "page_height": section.page_height,

            "left_margin": section.left_margin,

            "right_margin": section.right_margin,

            "top_margin": section.top_margin,

            "bottom_margin": section.bottom_margin,

            "header_distance": section.header_distance,

            "footer_distance": section.footer_distance

        })

    return sections

# --------------------------------------------------
# Comments
# --------------------------------------------------

def extract_comments(file_path):
    """
    Extract Word comments from comments.xml.
    """

    comments = []

    try:

        with ZipFile(file_path) as docx:

            if "word/comments.xml" not in docx.namelist():
                return comments

            xml = docx.read("word/comments.xml")

            root = etree.fromstring(xml)

            namespaces = {
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            }

            for comment in root.findall("w:comment", namespaces):

                comment_id = comment.get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id"
                )

                author = comment.get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}author"
                )

                text = ""

                for t in comment.findall(".//w:t", namespaces):

                    if t.text:
                        text += t.text

                comments.append({

                    "comment_id": comment_id,

                    "author": author,

                    "text": text

                })

    except Exception as e:

        comments.append({

            "error": str(e)

        })

    return comments

# --------------------------------------------------
# Bookmarks
# --------------------------------------------------

def extract_bookmarks(file_path):
    """
    Extract bookmarks from a Word document.
    """

    bookmarks = []

    try:

        with ZipFile(file_path) as docx:

            if "word/document.xml" not in docx.namelist():
                return bookmarks

            xml = docx.read("word/document.xml")

            root = etree.fromstring(xml)

            ns = {
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            }

            for bookmark in root.findall(".//w:bookmarkStart", ns):

                bookmarks.append({

                    "bookmark_id": bookmark.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id"
                    ),

                    "bookmark_name": bookmark.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name"
                    )

                })

    except Exception as e:

        bookmarks.append({

            "error": str(e)

        })

    return bookmarks

# --------------------------------------------------
# Fields
# --------------------------------------------------

def extract_fields(file_path):
    """
    Extract Word field codes from document.xml.
    """

    fields = []

    try:

        with ZipFile(file_path) as docx:

            if "word/document.xml" not in docx.namelist():
                return fields

            xml = docx.read("word/document.xml")

            root = etree.fromstring(xml)

            ns = {
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            }

            field_number = 1

            # Complex field instructions
            for instr in root.findall(".//w:instrText", ns):

                if instr.text:

                    field_code = instr.text.strip()

                    field_type = field_code.split()[0] if field_code else ""

                    fields.append({

                        "field_number": field_number,

                        "field_type": field_type,

                        "field_code": field_code

                    })

                    field_number += 1

    except Exception as e:

        fields.append({

            "error": str(e)

        })

    return fields

# --------------------------------------------------
# Content Controls (Structured Document Tags)
# --------------------------------------------------

def extract_content_controls(file_path):
    """
    Extract Structured Document Tags (Content Controls)
    from a Word document.
    """

    content_controls = []

    try:

        with ZipFile(file_path) as docx:

            if "word/document.xml" not in docx.namelist():
                return content_controls

            xml = docx.read("word/document.xml")

            root = etree.fromstring(xml)

            ns = {
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            }

            control_number = 1

            for sdt in root.findall(".//w:sdt", ns):

                alias = ""

                tag = ""

                text = ""

                alias_node = sdt.find(".//w:alias", ns)

                if alias_node is not None:
                    alias = alias_node.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val",
                        ""
                    )

                tag_node = sdt.find(".//w:tag", ns)

                if tag_node is not None:
                    tag = tag_node.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val",
                        ""
                    )

                for t in sdt.findall(".//w:t", ns):

                    if t.text:
                        text += t.text

                content_controls.append({

                    "control_number": control_number,

                    "alias": alias,

                    "tag": tag,

                    "text": text

                })

                control_number += 1

    except Exception as e:

        content_controls.append({

            "error": str(e)

        })

    return content_controls

# --------------------------------------------------
# Track Changes
# --------------------------------------------------

def extract_track_changes(file_path):
    """
    Extract tracked changes (insertions, deletions, moves)
    from a Word document.
    """

    track_changes = []

    try:

        with ZipFile(file_path) as docx:

            if "word/document.xml" not in docx.namelist():
                return track_changes

            xml = docx.read("word/document.xml")

            root = etree.fromstring(xml)

            ns = {
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            }

            change_number = 1

            # -------------------------
            # Insertions
            # -------------------------

            for ins in root.findall(".//w:ins", ns):

                text = ""

                for t in ins.findall(".//w:t", ns):

                    if t.text:
                        text += t.text

                track_changes.append({

                    "change_number": change_number,

                    "type": "Insertion",

                    "author": ins.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}author"
                    ),

                    "date": ins.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}date"
                    ),

                    "text": text

                })

                change_number += 1

            # -------------------------
            # Deletions
            # -------------------------

            for delete in root.findall(".//w:del", ns):

                text = ""

                for t in delete.findall(".//w:delText", ns):

                    if t.text:
                        text += t.text

                track_changes.append({

                    "change_number": change_number,

                    "type": "Deletion",

                    "author": delete.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}author"
                    ),

                    "date": delete.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}date"
                    ),

                    "text": text

                })

                change_number += 1

            # -------------------------
            # Move From
            # -------------------------

            for move in root.findall(".//w:moveFrom", ns):

                track_changes.append({

                    "change_number": change_number,

                    "type": "MoveFrom",

                    "author": move.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}author"
                    ),

                    "date": move.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}date"
                    )

                })

                change_number += 1

            # -------------------------
            # Move To
            # -------------------------

            for move in root.findall(".//w:moveTo", ns):

                track_changes.append({

                    "change_number": change_number,

                    "type": "MoveTo",

                    "author": move.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}author"
                    ),

                    "date": move.get(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}date"
                    )

                })

                change_number += 1

    except Exception as e:

        track_changes.append({

            "error": str(e)

        })

    return track_changes


# --------------------------------------------------
# Main Extractor
# --------------------------------------------------

def extract(file_path):

    document = Document(file_path)

    result = create_document_template(file_path)
    result["document"]["file_type"] = "word"
    result["metadata"] = extract_metadata(document)
    result["paragraphs"] = extract_paragraphs(document)
    result["headings"] = extract_headings(document)
    result["tables"] = extract_tables(document)
    result["images"] = extract_images(document, file_path)
    result["hyperlinks"] = extract_hyperlinks(document)
    result["headers"] = extract_headers(document)
    result["footers"] = extract_footers(document)
    result["sections"] = extract_sections(document)
    result["comments"] = extract_comments(file_path)
    result["bookmarks"] = extract_bookmarks(file_path)
    result["fields"] = extract_fields(file_path)
    result["content_controls"] = extract_content_controls(file_path)
    result["track_changes"] = extract_track_changes(file_path)

    return result
